import json, math
from pathlib import Path
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

out=Path(__file__).resolve().parent/'report';out.mkdir(exist_ok=True)
charts=out/'charts';charts.mkdir(exist_ok=True)
raw=json.loads((Path(__file__).resolve().parent/'data/validated-monthly-analysis.json').read_text(encoding='utf8'))
season=json.loads((Path(__file__).resolve().parent/'data/seasonal-booking-estimate.json').read_text(encoding='utf8'))
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'axes.spines.top':False,'axes.spines.right':False,'figure.dpi':170})
def save(name):
    titles={'airport.png':'BGY annual passenger movements 2019–2025','airport-city.png':'BGY and City of Bergamo growth since 2019','growth.png':'City of Bergamo guest nights by year','season.png':'City of Bergamo monthly guest nights in 2024','allocation.png':'Seasonal allocation of 180 scenario booked nights','bridge.png':'Short-term rental scenario annual cash bridge','sensitivity.png':'Short-term rental price and booking sensitivity','setup.png':'Booked nights required for five-year cost recovery'}
    plt.gca().set_title(titles[name],fontsize=11)
    plt.tight_layout();plt.savefig(charts/name,bbox_inches='tight');plt.close()
ys=list(range(2019,2025));vals=[sum(x['guest_nights'] for x in raw if x['year']==y) for y in ys]
fig,ax=plt.subplots(figsize=(7,3));ax.plot(ys,[v/1e6 for v in vals],marker='o',label='Regional series through 2024');ax.scatter([2025],[1.359923],marker='D',label='2025 local report');ax.set_ylabel('Guest nights (millions)');ax.set_xticks(ys+[2025]);ax.set_ylim(0,1.55);ax.legend(frameon=False);save('growth.png')
mon=['Jan','Feb','Mar','Apr','May','Jun','Jul','Aug','Sep','Oct','Nov','Dec']
r24=sorted([x for x in raw if x['year']==2024],key=lambda x:x['month'])
fig,ax=plt.subplots(figsize=(7,3));ax.bar(mon,[x['guest_nights']/1000 for x in r24],color='#1C4C74');ax.set_ylabel('Observed guest nights (thousands)');save('season.png')
fig,ax=plt.subplots(figsize=(7,3));ax.bar(mon,[x['bookings']['180'] for x in season['months']],color='#1C4C74');ax.set_ylabel('Scenario booked nights');ax.set_ylim(0,22);save('allocation.png')
fig,ax=plt.subplots(figsize=(7,2.5));labels=['Gross','Tax','Platform','Cleaning','Variable','Fixed','Net cash'];v=[20700,-4347,-3208.5,-1800,-1800,-3000,6544.5];running=0
for i,value in enumerate(v):
    if i in [0,6]:ax.bar(i,value,color='#1C4C74' if i==0 else '#32B8B2');top=value
    else:ax.bar(i,-value,bottom=running+value,color='#A8B7C5');top=running
    ax.text(i,top+350,f'{abs(value):,.0f}',ha='center',fontsize=9)
    if i!=6:running+=value
ax.set_xticks(range(7),labels);ax.set_ylabel('EUR per year');ax.set_ylim(0,23500);save('bridge.png')
fig,ax=plt.subplots(figsize=(7,3));ns=[150,180,210]
for a,c in [(100,'#64798B'),(115,'#1C4C74'),(130,'#32B8B2')]:ax.plot(ns,[n*(a*.635-20)-3000 for n in ns],marker='o',label=f'Nightly revenue EUR {a}',color=c)
ax.axhline(6732,color='#A34A32',linestyle='--',label='Long-term scenario EUR 6,732');ax.set_xticks(ns);ax.set_xlabel('Annual booked nights');ax.set_ylabel('Annual net cash income EUR');ax.legend(frameon=False,fontsize=8);save('sensitivity.png')
fig,ax=plt.subplots(figsize=(7,2.8));setup=[0,5000,10000,15000];n=[math.ceil((6732+3000+s/5)/(115*.635-20)) for s in setup];ax.bar([f'{s:,}' for s in setup],n,color='#1C4C74')
for i,v in enumerate(n):ax.text(i,v+4,str(v),ha='center')
ax.set_ylim(0,280);ax.set_xlabel('Incremental preparation cost EUR');ax.set_ylabel('Required booked nights');save('setup.png')

airport=json.loads((Path(__file__).resolve().parent/'data/airport-city-annual.json').read_text())
fig,ax=plt.subplots(figsize=(7,2.8));ax.bar([r['year'] for r in airport],[r['airport_passengers']/1e6 for r in airport],color='#1C4C74');ax.set_xticks(range(2019,2026));ax.set_ylabel('Annual passenger movements (millions)');save('airport.png')
fig,ax=plt.subplots(figsize=(7,3));
for key,label,col in [('airport_passengers','BGY passengers','#64798B'),('city_arrivals','City arrivals','#1C4C74'),('city_guest_nights','City guest nights','#32B8B2')]:
    ax.plot([r['year'] for r in airport],[100*r[key]/airport[0][key] for r in airport],marker='o',label=label,color=col)
ax.axhline(100,color='#BBC5CC',linestyle=':',linewidth=1);ax.set_xticks(range(2019,2026));ax.set_ylabel('Index, 2019 = 100');ax.legend(frameon=False,fontsize=8);save('airport-city.png')

d=Document();s=d.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11);s.top_margin=s.bottom_margin=Inches(.7);s.left_margin=s.right_margin=Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=d.styles[name];st.font.name='Calibri';st.font.size=Pt(11 if name=='Normal' else 22 if name=='Title' else 15 if name=='Heading 1' else 12);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.08
    for el in list(st.element.iter(qn('w:pBdr'))):el.getparent().remove(el)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.add_run('Pedram Kojouri | Research edition | ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)

header=s.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
hl=OxmlElement('w:hyperlink');hl.set(qn('r:id'),header.part.relate_to('https://www.linkedin.com/in/pedram-kojouri/',RT.HYPERLINK,is_external=True))
hr=OxmlElement('w:r');prop=OxmlElement('w:rPr');co=OxmlElement('w:color');co.set(qn('w:val'),'1C4C74');prop.append(co);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');prop.append(sz);hr.append(prop);ht=OxmlElement('w:t');ht.text='Pedram Kojouri  ·  linkedin.com/in/pedram-kojouri';hr.append(ht);hl.append(hr);header._p.append(hl)
s.header_distance=Inches(.25)

pending_break=False
def p(t,style=None):
    global pending_break
    x=d.add_paragraph(t,style)
    if pending_break:
        x.paragraph_format.page_break_before=True;pending_break=False
    return x
def h(t):p(t,'Heading 1')
def sub(t):p(t,'Heading 2')
def page():
    global pending_break
    pending_break=True
def table(headers,rows):
    t=d.add_table(rows=1,cols=len(headers))
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for ri,row in enumerate(t.rows):
        for cell in row.cells:
            pr=cell._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
            pr.append(b)
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E9EFF4' if ri==0 else 'FFFFFF' if ri%2 else 'F6F8FA');pr.append(sh)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after=Pt(4)
                for run in paragraph.runs:run.font.size=Pt(9)
    p('')
def fig(name,caption,explanation):
    d.add_picture(str(charts/name),width=Inches(6.6));p(caption,'Caption');p(explanation)
p('Rental Strategies in Bergamo: More Tourists, More Profit?','Title');p('A comparison of short-term and long-term rentals');p('Pedram Kojouri | Complete research report | 13 October 2026')
sub('Executive finding')
p('Bergamo has a growing accommodation market, but tourism growth does not establish a general net-income advantage for short-term renting. In the illustrative owner-operated scenario, short-term rental net cash income is EUR 6,544.50 per year, compared with EUR 6,732 for the long-term alternative. The EUR 187.50 gap is too small, relative to uncertainty in the inputs, to support a decisive ranking.')
p('The market analysis concerns the City of Bergamo. Financial calculations explain conditions under which strategies may differ; they do not estimate average city-wide profit. Acquisition, financing, professional management and incremental preparation are outside the base calculation. No monetary deduction is made for the owner’s time.')
page();p('Contents','Title')
tocp=p('');toc=OxmlElement('w:fldSimple');toc.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');tocp._p.append(toc)
page();h('1 Introduction and analytical scope')
p('The project began with an apparent opportunity: more visitors might mean more rental income. Historic Città Alta and access to BGY provided a reason to investigate, but were not treated as proof of profitability. The research question became: under what conditions would short-term renting retain more annual cash than a residential lease?')
p('The investigation moved through three questions. Is accommodation demand growing? Is rental performance improving after allowing for available nights and competing supply? How much revenue remains after tax and operating expenses? The last question introduces long-term renting as an opportunity cost rather than comparing a nightly price directly with monthly rent.')
sub('Geography and evidence')
p('Comune di Bergamo, ISTAT code 016024, is the principal unit. Province-wide totals and municipalities near the airport are context only. Airport passengers are not unique overnight visitors to the city. Events and holidays may explain variations, but no event revenue premium or airport-to-booking conversion is assumed.')
p('Official regional data support the historical city series; local official reports supplement annual context. FULL research supplies historical Airbnb estimates derived from AirDNA. Asking-rent listings illustrate alternatives, not executed contract prices. FULL and AirDNA are not independent corroborating sources. Scenario calculations are kept separate from observations.')
sub('Definitions and validation')
table(['Measure','Interpretation'],[['Arrivals','Recorded guest arrivals; not necessarily unique people.'],['Guest nights','One guest for one night. Three people for two nights produce six guest nights but two apartment booked nights.'],['Booked nights','Nights reserved for a listing.'],['Occupancy','Booked nights divided by available nights, not automatically by 365.'],['ADR','Provider-defined revenue per booked night; included cleaning charges and fee deductions must be checked.'],['Net cash income','Revenue less modelled cash costs and tax, not total investment return.']])
p('Monthly observations for 2019–2024 were filtered to the city, checked for twelve distinct months and reconciled to the regional annual dataset. All six years reconcile within that source. Reference years and geography remain explicit. Financial formulas are checked separately from market interpretation. [1–2]')
page();h('2 Tourism demand and competing supply')
p('Regional city data show 791,353 guest nights in 2019 and 1,208,189 in 2024, approximately 52.7% growth. Arrivals rose from 414,871 to 601,824. The 2025 official local report records 666,254 city arrivals and 1,359,923 guest nights. This indicates growth beyond the pre-pandemic level, although the 2025 total comes from a separate release. [1–4]')
fig('growth.png','Figure 1  City guest nights from 2019 to 2025','The line uses the reconciled regional series through 2024. The separate diamond identifies the 2025 local-report total. Growth measures aggregate guest demand, not income or profit per listing. Sources [1–4].')
sub('Source reconciliation')
p('The May 2025 local release gives 601,825 arrivals and 1,208,193 guest nights for 2024, compared with 601,824 and 1,208,189 in the regional extract. October differs by 27 guest nights: 123,433 versus 123,460. The cause is not confirmed. The report uses the regional series consistently for historical calculations rather than silently blending releases. [1–3]')
sub('Supply and geography')
p('The Province recorded 3,359,760 guest nights in 2025, approximately 38% above 2019. Its non-hotel establishment count increased from 3,443 to 4,405 between 2024 and 2025. These counts cannot establish a shortage or saturation in the city: the geography differs and establishments are not beds or open days. National coverage of private non-business rentals has changed; the effect on the local series still needs validation. [4, 6]')
p('FULL reports 1,884 city Airbnb listings active for at least one day in 2024, up 74.3% from 2017. Booked nights rose 182.8% over that period. Availability, listing mix, prices and costs prevent a direct conclusion that every host’s return improved. [5, pp. 12–18]')
page();sub('2.1 The role of Bergamo Airport in tourism demand')
p('BGY adds a large regional access channel to the business story. We joined Assaeroporti annual passenger totals to city arrivals and guest nights by calendar year for 2019–2025. For each year, the January–December passenger total was selected, rather than the December-only row. The annual summary and passenger-detail tables agree. City values through 2024 use the reconciled regional extract; 2025 uses the separately identified provincial publication. [1–4, 37–43]')
table(['Year','BGY passengers','City arrivals','City guest nights'],[[str(r['year']),f"{r['airport_passengers']:,}",f"{r['city_arrivals']:,}",f"{r['city_guest_nights']:,}"] for r in airport])
fig('airport.png','Figure 2  BGY annual passenger movements','Source: Assaeroporti, January–December cumulative passenger totals, 2019–2025. [37–43] This airport activity metric counts passenger movements rather than unique tourists staying in Bergamo. The 2020–2021 decline and later recovery affect comparisons across the full window.')
page();sub('2.2 Airport and city growth compared')
fig('airport-city.png','Figure 3  Indexed annual airport and city trends','Each series is divided by its 2019 value and multiplied by 100. Indexing compares relative growth across different units without treating passengers, arrivals and guest nights as equivalent counts. Sources: [1–4, 37–43]. The 2025 city value comes from a separate release.')
table(['Year-on-year change','BGY passengers','City arrivals','City guest nights'],[[str(airport[i]['year'])]+[f"{(airport[i][k]/airport[i-1][k]-1)*100:+.1f}%" for k in ['airport_passengers','city_arrivals','city_guest_nights']] for i in range(1,len(airport))])
p('Relative to 2019, BGY passengers were 25.2% higher in 2024 and 22.2% higher in 2025. City guest nights grew by 52.7% and 71.8% respectively. In 2025, airport passengers fell 2.4%, while city arrivals grew 10.7% and guest nights grew 12.6%. The city therefore continued expanding even when airport activity declined. The series share a recovery period but do not move in lockstep.')
p('This is a descriptive comparison of seven annual observations. Pandemic disruption, accommodation supply, event calendars and wider travel trends can move both series. We do not estimate a causal effect or a passenger-to-stay conversion rate. Airport passengers include departing and arriving trips and travellers whose destination may lie outside the city. Accommodation also serves visitors using other transport. Dividing city arrivals by airport passengers would not identify the share staying locally.')
p('Business implication: airport access deserves attention when considering guest segments, arrival times and average stay length, but the airport trend does not justify increasing the 180-night assumption or adding a location premium. One-night airport stays could increase cleaning turnovers; this remains a hypothesis. Neighbourhood effects, airline-specific traffic and net profit attributable to BGY require matched property and guest-origin evidence.')

page();h('3 Short-term performance and seasonality')
p('FULL reports overall Airbnb occupancy falling from 57.8% in 2023 to 54.0% in 2024, a reduction of 3.8 percentage points. Entire apartments in 2024 show 54.9% occupancy and approximately EUR 115 ADR. Across all listing types, average annual bookings were 134 nights and average revenue EUR 14,182. The all-listing figure is not an entire-apartment benchmark. [5, pp. 16–20]')
p('The financial model uses EUR 115 as illustrative gross nightly revenue before the assumed host fee and income tax. It is not a validated direct import of provider ADR. AirDNA includes cleaning charges and may adjust service fees; revenue definitions must be reconciled to avoid subtracting a fee twice. Multiplying separately averaged price and occupancy also need not reproduce average listing revenue. [7–8]')
fig('season.png','Figure 4  Observed monthly city guest nights in 2024','July leads arrivals; October leads guest nights. The bars include hotels and non-hotels. They identify the city’s seasonal demand, not Airbnb occupancy or fully booked apartments. Source [1].')
p('April–October accounts for 65.5% of city guest nights in 2019, 67.3% in 2022, 67.4% in 2023 and 67.3% in 2024. October leads guest nights in each of those years. High demand does not mean full occupancy: FULL reports a 2024 overall monthly maximum of 69% in May and a minimum of 34% in January. [1, 5]')
page();sub('Monthly allocation and its limits')
p('For each reference year, monthly daily guest nights are divided by that year’s annual daily mean. The indices are averaged equally across 2019, 2022, 2023 and 2024, multiplied by non-leap-year month lengths and normalised into shares. Largest-remainder rounding preserves the annual scenario total. The calendar calculation is not observed 2025 data.')
p('Annual totals of 150, 180 and 210 booked nights are sensitivity assumptions, not confidence intervals. The 180-night base implies 54.5% occupancy only if 330 nights are available. Availability and annual bookings remain assumptions. The method estimates distribution, not the correct annual total.')
fig('allocation.png','Figure 5  Monthly allocation of the 180-night scenario','January receives nine nights; July, September and October receive eighteen each. April–October receives 120 nights or 66.7%. These are allocations, not reservations observed on Airbnb. Tourism seasonality is an unvalidated proxy for short-term seasonality. Derived from [1].')
p('With a constant price and constant per-booked-night costs, changing the monthly distribution does not change annual net cash. The allocation improves the description of timing; seasonal prices and stay lengths would be needed to measure its financial effect. Allocating annual fixed costs evenly across months is reporting, not actual payment timing.')
sub('Neighborhood differences')
table(['Area','Reported annual revenue per listing in 2024'],[['Centro–Papa Giovanni','EUR 21,829'],['Città Alta','EUR 19,712'],['Centro–Pignolo','EUR 17,321'],['Centro–Sant’Alessandro','EUR 14,976'],['Borgo Palazzo','EUR 14,740']])
p('These averages are unadjusted descriptive comparisons, not a profitability ranking. Size, capacity, quality, available days and acquisition costs may differ. No causal location premium is estimated. A table is used instead of a league-chart presentation to avoid implying a superior investment area. [5, pp. 28–32]')
page();h('4 The long-term rental comparison')
p('A defensible alternative must match location, size, furnishing, condition and contract terms. The present listing sample is too small for a city-wide median. EUR 900 monthly rent is an asking-rent-supported scenario anchor, not a verified city benchmark.')
table(['Listing and size','Base monthly rent','Contract and separate charges'],[['Calepio; Borgo Palazzo; 47 m²','EUR 900','Furnished; 4+4 lease; charges not stated.'],['Borfuro 9; central; 55 m²','EUR 1,000','Annual renewable; EUR 110 condominium charges separately.'],['Moroni; near Piazza Pontida; 50 m²','EUR 700','Temporary 12 months; EUR 50 charges separately.']])
p('These contract classes are not pooled. Separate charges are not owner rent, and asking prices are not verified signed rents. An advertiser’s refusal of B&B use is an offer condition, not proof of a neighborhood-wide legal ban. Listings that could not be re-opened are not counted as confirmed active observations. [10–12]')
p('Immobiliare.it reports September 2026 residential asking rents of EUR 13.57 per m² per month for the city, EUR 14.31 for combined Borgo Palazzo/Centro/Stazione and EUR 14.12 for Città Alta. These broad residential averages are not a furnished one-bedroom valuation. [9]')
sub('Long-term cash assumptions')
table(['Item','Annual scenario'],[['Monthly rent and paid months','EUR 900; 12 months'],['Gross revenue','EUR 10,800'],['Assumed eligible tax at 21%','EUR 2,268'],['Owner cash-cost allowance','EUR 1,800'],['Net cash income','EUR 6,732']])
p('Twelve paid months is a contract-income scenario, not observed zero vacancy or arrears. Eleven paid months produces EUR 6,021 using the same allowances. The EUR 1,800 is unverified; owner maintenance, insurance, property taxes and letting fees need confirmation. The study does not estimate average executed city rents or losses. Historical short-term performance and 2026 asking rents cover different periods.')
page();sub('Expanded rental listing screening')
p('A purposive portal review extends the original three observations with eighteen additional candidate listings. The registry contains twenty-one unique listing IDs, not a representative city sample. Screening retains six furnished one-bedroom offers with a 4+4 option and approximately 45–60 dwelling square metres in the primary cohort. Other durations, serviced accommodation, conflicting prices and insufficient detail remain separate or excluded. Listings are snapshots observed on 13 October 2026; a subsequent failed fetch is recorded rather than treated as proof that an offer remains active.')
table(['Primary offer / area','Size m²','Rent EUR/month','Terms or comparability flag'],[['Ronzoni / San Paolo','60 dwelling','1,400','4+4; garage bundled; larger commercial area'],['Sant Ambrogio / Canovine','53','850','4+4; optional garage excluded'],['Locatelli / Centro','60','1,200','4+4'],['San Fermo / Borgo Palazzo','60 advertised','900','4+4; later re-fetch failed'],['Ruspini / Canovine','60','750','4+4 or transitory offered'],['Calepio / Borgo Palazzo','47','900','4+4']])
p('The within-cohort asking-rent median is EUR 900, with a range of EUR 750–1,400. This is a descriptive sample statistic, not a city rent estimate, and the six offers are not perfectly matched on parking, floor, quality or surface definition. Removing the garage-bundled Ronzoni offer leaves the median at EUR 900. EUR 900 is retained as a scenario anchor, with EUR 700–1,100 sensitivity rather than a claim of precision.')
p('The suggested target of 15–30 genuinely comparable offers has not been met. Adding serviced or uncertain offers merely to reach that number would weaken the analysis. Contract and parking effects cannot be estimated from this cohort. The public registry provides URLs, classification and exclusion reasons without photographs, contacts or copied listing descriptions.')

page();h('5 Revenue, Costs and Net Cash Income')
p('The base rental is owner-operated and already suitable for guests. Cleaning is entirely owner-paid, with no cleaning charge collected separately. No hypothetical cost for the owner’s time is deducted. Professional management would be an additional real payment.')
table(['Input','Base value and status'],[['Bookings and availability','180 / 330 nights; scenario assumptions'],['Gross nightly revenue','EUR 115; reference-informed assumption'],['Host fee and eligible income tax','15.5% and 21%; scenario rates'],['Cleaning and stay length','EUR 30 per turnover; 3 nights; unverified'],['Other variable costs','EUR 10 per booked night; allowance'],['Annual fixed costs','EUR 3,000: common EUR 1,800 + STR-specific EUR 1,200']])
p('EUR 115 is a reference-informed modelling assumption for gross accommodation revenue before host fees and income tax. It has not been verified as an observed, directly comparable gross ADR. The scenario collects no separate cleaning charge; provider ADR definitions must be reconciled before replacing this input.')
p('At three nights per stay, 180 nights imply sixty turnovers and EUR 1,800 cleaning expense. At two nights, the same cleaning price costs EUR 2,700; at five nights, EUR 1,080. Laundry, linen and materials must be assigned once, not duplicated across allowances.')
p('The 15.5% fee is a host-paid structure assumption, not universal. Airbnb documents different structures and says displayed service fees include VAT where applicable; VAT is not added again here. Account and channel mix require verification. [13]')
p('The 21% tax is conditional on eligibility. Equal model rates do not prove equal legal treatment. Tax is applied to model gross revenue before platform fees. Separately collected and remitted tourist tax is not retained revenue. IMU, TARI and property-specific compliance expenses are not individually priced and must be checked against the allowances. [14–16]')
fig('bridge.png','Figure 6  Short-term rental cash bridge','EUR 20,700 gross is reduced by EUR 4,347 tax, EUR 3,208.50 host fees, EUR 1,800 cleaning, EUR 1,800 variable costs and EUR 3,000 fixed costs. EUR 6,544.50 remains, about 31.6% of gross. This is a scenario margin, not an observed city margin. Source: project assumptions and calculations; EUR 115 is not a verified comparable gross ADR.')
page();sub('Fixed-cost budget and uncertainty')
p('Net cash = booked nights × gross nightly revenue × 0.635 − booked nights × EUR 20 − EUR 3,000. The EUR 20 combines cleaning and other variable costs. Acquisition, debt service, incremental preparation and professional management are excluded; exceptional maintenance may exceed allowances.')
p('The following is an illustrative allocation of the existing EUR 3,000 allowance, not local supplier quotes or measured owner spending. Ranges are analyst-selected stress values, not confidence intervals. Property tax eligibility, tenant-paid charges and actual bills must replace them for a property decision. Keeping the base total unchanged isolates the effect of transparency rather than silently changing the result.')
table(['Annual item','Base EUR','Stress range EUR'],[['Common insurance','250','150–400'],['Common routine maintenance','400','200–700'],['Common property-tax reserve','800','0–1,300'],['Common owner building charges','350','150–600'],['Common subtotal','1,800','900 / 1,800 / 2,700 tested as totals'],['STR standing utility / internet charges','400','200–650'],['STR administration / compliance reserve','300','100–500'],['STR guest-wear replacement reserve','400','200–700'],['STR marketing / miscellaneous','100','0–150'],['STR incremental subtotal','1,200','600 / 1,200 / 1,800 tested as totals']])
p('Standing charges exclude usage already charged at EUR 10 per booked night. Guest-wear replacement excludes ordinary maintenance in the common allowance. Cleaning, laundry and consumables must be assigned once within cleaning or variable expenses. A reserve is a budgeting allowance, not an observed payment; realised annual cash will differ.')
p('Changing a truly common cost by the same amount under both strategies changes both net incomes equally and leaves their difference unchanged. STR-specific expenses change the ranking. Some charges may not be common in practice; transferability must be checked, particularly building charges, waste tax and insurance.')

page();sub('Regulatory feasibility')
p('Classification, regional and municipal filings, CIN, guest reporting and safety must be checked. CIN does not establish full local compliance. Reviewed municipal provisions address building usability and installations, with additional conditions in mapped historic areas. Applicable amendments and later court decisions remain to be confirmed. Unresolved compliance costs are not silently assumed to be zero. [16–18]')
h('6 Sensitivity and preparation costs')
p('Short-term base net cash is EUR 6,544.50 versus EUR 6,732 long-term. The EUR 187.50 difference is about 2.8% of long-term net cash. It is a mechanical model result, not a robust preference when input uncertainty is unquantified.')
fig('sensitivity.png','Figure 7  Booking and nightly-price sensitivity','The dashed line is the long-term scenario, not a market average. At 180 nights, short-term net cash ranges from EUR 4,830 at EUR 100 to EUR 8,259 at EUR 130. Higher prices may reduce reservations. The combinations are conditional calculations, not forecasts. Source: project scenario calculations.')
sub('Break-even and calculation checks')
p('At EUR 115, each booked night contributes EUR 53.025 before fixed expenses. Matching long-term cash requires (6,732 + 3,000) / 53.025 = 183.54 nights, rounded up to 184. At 180 nights, the required price is EUR 116.64. Increasing price by EUR 1 adds EUR 114.30 annual cash, holding bookings fixed. One additional percentage point of host fee costs EUR 207; EUR 5 less cleaning per turnover saves EUR 300 across sixty turnovers.')
page();sub('Long-term rent and differential-cost sensitivity')
table(['Monthly rent EUR','Long-term net EUR/year','STR minus LTR EUR/year'],[[f'{r:,}',f'{r*12*.79-1800:,.2f}',f'{6544.5-(r*12*.79-1800):+,.2f}'] for r in [700,800,900,1000,1100]])
p('All rows assume twelve paid months, eligible 21% tax and EUR 1,800 owner costs. These are scenario values, not five observed market segments. At the base short-term performance, equal cash occurs at about EUR 880.22 monthly long-term rent. The choice of alternative rent therefore matters as much as the short-term inputs.')
table(['STR additional fixed EUR/year','LTR 800/month','LTR 900/month','LTR 1,000/month'],[[str(c)]+[f'{180*(115*.635-20)-1800-c-(r*12*.79-1800):+,.2f}' for r in [800,900,1000]] for c in [600,1200,1800]])
p('Cells show short-term minus long-term annual net cash, in EUR. Positive cells favour short-term on cash alone; negative cells favour long-term. Common costs are held at EUR 1,800 and all other base inputs are held constant. At EUR 900 rent, changing the STR-specific budget from EUR 1,200 to EUR 600 reverses the gap from minus EUR 187.50 to plus EUR 412.50. A EUR 600 higher budget produces minus EUR 787.50. Joint price and booking feasibility remains untested.')

page();sub('Additional preparation changes the threshold')
p('Only expenditure above the long-term alternative is relevant to conversion. A five-year simple recovery target adds preparation cost divided by five to the required annual surplus. It leaves operating cash unchanged and is not a market-required horizon or a discounted investment appraisal.')
fig('setup.png','Figure 8  Booking threshold with five-year simple recovery','At EUR 115 nightly revenue, the model requires 184, 203, 222 and 241 nights for EUR 0, EUR 5,000, EUR 10,000 and EUR 15,000 additional preparation. The latter thresholds match long-term cash and recover the cost difference over five years. They do not predict achievable occupancy. Source: project calculations with explicit scenario inputs.')
table(['Additional preparation','Nights at EUR 115','Price at 180 nights'],[['EUR 0','184','EUR 116.64'],['EUR 5,000','203','EUR 125.39'],['EUR 10,000','222','EUR 134.14'],['EUR 15,000','241','EUR 142.89']])
p('Required booked nights are always rounded UP to the next whole night using the ceiling function; conventional rounding would understate the minimum. Unrounded thresholds are 183.54, 202.40, 221.25 and 240.11 nights using the stated formula. These are scenario requirements, not occupancy forecasts.')
p('The two performance columns are alternatives, not simultaneous improvements. Financing, inflation, discounting and residual equipment value are excluded. Already suitable properties face a close operating comparison; costly conversion requires a substantiated surplus. More visitors alone cannot justify that investment.')
page();h('7 Research methodology and challenges')
p('The investigation began with an opportunity hypothesis, then changed its emphasis as evidence challenged the assumption that more visitors necessarily improve owner profit. The workflow combines validated market observations with an explicit decision model. It stops short of identifying a causal airport effect or an average investment return.')
table(['Research challenge','Action and remaining limit'],[['City versus province','Filter tourism to Bergamo 016024. Province supply and airport traffic remain contextual, not city apartment reservations.'],['Conflicting releases','Reconcile monthly and annual regional totals, retain local-release differences including 27 October guest nights. Revision cause unresolved.'],['Guest nights versus booked nights','Use tourism for seasonal shape. Do not convert peak months into full occupancy. Annual 180 bookings and 330 availability remain assumptions.'],['Price definitions','Treat EUR 115 as a reference-informed gross model input. Provider ADR may include cleaning or fee adjustments.'],['Rental comparability','Screen 21 candidate IDs; retain six furnished 4+4-option offers. Separate temporary and serviced units. Target 15–30 comparable offers remains unmet.'],['Operating budgets','Itemise allowances and vary differential costs. Values are illustrative budgets, not supplier quotes. Avoid double counting cleaning and usage.'],['Airport aggregation','Join seven annual totals by year; check cumulative versus December-only records. Compare indexed trends and annual growth without claiming causality.'],['Different time windows','Identify historical 2024 short-term estimates and 2026 rental asks. No same-period matched profitability panel exists.'],['Decision uncertainty','Stress-test both strategies and preparation costs. Do not interpret sensitivity ranges as confidence intervals.']])
sub('Reproducibility and review')
p('The public repository includes frozen numerical inputs, listing screening decisions and executable Python calculations. Monthly sums match six annual controls. Seasonal allocation preserves annual totals. Break-even checks verify that the chosen whole-night threshold meets the target and one fewer night does not. Airport totals reconcile across two tables within each source year. Financial and airport comparison scripts run offline.')
p('Future evidence that could change the ranking includes comparable realised price and booked nights, contract rents for the same units, supplier quotes and applicable legal costs. A larger matched dataset would permit property-level modelling. Seven annual aggregates alone do not justify causal inference or a reliable multivariable regression.')

page();h('8 Conclusions and limitations')
p('More tourists do not necessarily mean more profit. City demand and Airbnb activity have expanded, but this study does not establish a general net-income advantage for short-term renting across Bergamo. Market opportunity and an owner’s retained cash are different outcomes.')
p('The base scenarios are close: EUR 6,544.50 short-term versus EUR 6,732 long-term. Stronger bookings, higher realised prices or lower expenses can reverse the result. Their joint attainability needs evidence. Preparation costs raise the requirement; buying and financing property remain a separate question.')
sub('Supported findings')
p('The city has substantial accommodation demand and a recurring broad high-demand season. Seasonal allocation should replace a uniform description of demand. Taxes and operating expenses reduce the attraction implied by gross short-term receipts. The financial choice depends on the rental alternative, cost structure and conversion expenditure.')
sub('Unresolved evidence')
p('The airport comparison establishes relative trends and a 2025 divergence, but does not identify local stays or a causal profit effect. The study has not measured typical city profit, optimal investment neighborhoods, a guaranteed 180-night total, realised monthly ADR, or the return on purchased property. Monthly tourism is a proxy for reservations; 330 nights of supply remains assumed. The official monthly dataset extracted here ends in 2024; a complete 2025–2026 monthly city series is not validated.')
p('A matched sample needs size, capacity, furnishing, condition, neighborhood and available nights. Long-term rents need more observations, deduplication and contract grouping. Provider revenue must be reconciled before fee subtraction. Monthly price and reservations should be paired within the same segment and period; unrelated averages do not support a correlation analysis.')
p('Utilities, cleaning, maintenance, insurance, property taxes and compliance require documented costs. Vacancy and arrears must be stress-tested. Data releases differ slightly and revisions must be retained. Owner time is excluded by design, but management effort remains a qualitative difference.')
p('Final answer: short-term renting is a credible business opportunity in Bergamo, but the present evidence supports conditional competitiveness, not a general recommendation to convert homes or a proven city-wide profit advantage. This is an edited market-informed scenario study pending comparable-property and cost validation.')
page();h('References')
p('Sources are placed at the end and cited by number in the text. Source checks and dataset retrieval took place during the project through 13 October 2026; not all pages were re-fetched for this editorial edition. Publication dates and coverage are distinguished. Asking-rent observations are snapshots, not a permanent live listing inventory.')
src=[
('1','Regione Lombardia / PoliS. Monthly municipal tourist flows, 2019–2024; Bergamo 016024. Twelve months reconciled for every year.','https://www.dati.lombardia.it/d/mzxz-sz25'),
('2','Regione Lombardia / PoliS. Annual municipal tourist flows. Total city series used as control.','https://www.dati.lombardia.it/d/vyxt-7jdx'),
('3','Provincia di Bergamo. Tourism 2024 release, 21 May 2025. City totals and October peak; small differences from the regional extract.','https://www.provincia.bergamo.it/cnvpbgrm/po/mostra_news.php?id=2163'),
('4','Provincia di Bergamo. Tourism 2025 and early 2026 release, 14 May 2026. City/province separately; attached report pp. 8 and 15–21. Early 2026 figures labelled estimated.','https://www.provincia.bergamo.it/cnvpbgrm/po/mostra_news.php?area=H&id=2591'),
('5','Chiodelli F. and Cassiani F. (2026). Chi gestisce Airbnb a Bergamo e dintorni. FULL. Author-posted research using AirDNA; data 2017–2024. Relevant pp. 12–20 and 28–32.','https://www.researchgate.net/publication/405451146_Chi_gestisce_Airbnb_a_Bergamo_e_dintorni_L%27evoluzione_del_fenomeno_degli_affitti_brevi_tra_il_2017_e_il_2024'),
('6','ISPRA. Tourism infrastructure statistics and notes on private non-business rental coverage.','https://indicatoriambientali.isprambiente.it/it/infrastrutture-turistiche'),
('7','AirDNA Help Center. ADR definition, including cleaning charges.','https://help.airdna.co/en/articles/8062173-how-does-airdna-calculate-average-daily-rate-adr'),
('8','AirDNA Help Center. Data updates, fee treatment and calendar classification.','https://help.airdna.co/en/articles/9305810-airdna-data-updates'),
('9','Immobiliare.it. Bergamo residential asking-rent overview; September 2026 reference period.','https://www.immobiliare.it/mercato-immobiliare/lombardia/bergamo/'),
('10','Idealista 36732340. Calepio 47 m²; EUR 900; furnished 4+4 lease.','https://www.idealista.it/immobile/36732340/'),
('11','Idealista 36084937. Borfuro 55 m²; EUR 1,000 plus EUR 110; annual contract.','https://www.idealista.it/immobile/36084937/'),
('12','Idealista 36830219. Moroni 50 m²; EUR 700 plus EUR 50; temporary 12-month lease.','https://www.idealista.it/immobile/36830219/'),
('13','Airbnb Help Center. Fee structures and VAT treatment.','https://www.airbnb.com/help/article/1857'),
('14','Agenzia delle Entrate. Building-income instructions; gross short-let receipts and eligible rate guidance. The 2026 return covers 2025 income.','https://infoprecompilata.agenziaentrate.gov.it/portale/sl/web/guest/semplificata-mod-fabbricati'),
('15','Regione Campania SURAP. Updated tax-guide notice, 16 April 2026; identifies 2026 threshold change.','https://surap.regione.campania.it/index.php/notizie/1096-locazioni-brevi-pubblicata-la-nuova-guida.html'),
('16','Ministero del Turismo. BDSR FAQ; CIN and safety, separate from local compliance.','https://www.ministeroturismo.gov.it/faq-banca-dati-strutture-ricettive-bdsr/'),
('17','Comune di Bergamo. Official mapping metadata for regulation 124/2025; not a complete legal clearance.','https://territorio.comune.bergamo.it/arcgis/rest/services/AtlanteBG/AppartamentiTuristici_Rif_PDR05_RegCom124_2025/MapServer'),
('18','Municipal regulation reproduced by Studio Buzzanca, December 2025; articles 1 and 4–8. Current official text and subsequent decisions require verification.','https://www.studiobuzzanca.com/wp-content/uploads/2025/12/regolamento-comunale-sui-requisiti-qualitativi-per-le-strutture-ricettive-non-alberghiere-.pdf')]
src.append((str(19),'Expanded listing 132189102; 4+4; garage bundled; 60 dwelling / 70.5 commercial; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/132189102/'))
src.append((str(20),'Expanded listing 133410334; Affitto, transitorio; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/133410334/'))
src.append((str(21),'Expanded listing 132644842; 4+4; optional garage excluded; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/132644842/'))
src.append((str(22),'Expanded listing 70888294; ; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/70888294/'))
src.append((str(23),'Expanded listing 80219117; Affitto, transitorio; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/80219117/'))
src.append((str(24),'Expanded listing 130803940; Affitto, 4+4; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/130803940/'))
src.append((str(25),'Expanded listing 132502222; Affitto, 4+4; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/132502222/'))
src.append((str(26),'Expanded listing 133336888; 4+4 offered; alternative temporary contract also offered; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/133336888/'))
src.append((str(27),'Expanded listing 129790224; Affitto; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/129790224/'))
src.append((str(28),'Expanded listing 127707676; Affitto, transitorio; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/127707676/'))
src.append((str(29),'Expanded listing 124638879; ; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/124638879/'))
src.append((str(30),'Expanded listing 133410286; ; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/133410286/'))
src.append((str(31),'Expanded listing 120078316; Affitto; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/120078316/'))
src.append((str(32),'Expanded listing 132219784; 4+4; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/132219784/'))
src.append((str(33),'Expanded listing 73676006; Affitto, libero; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/73676006/'))
src.append((str(34),'Expanded listing 121511946; ; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/121511946/'))
src.append((str(35),'Expanded listing 111380853; Affitto, transitorio; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/111380853/'))
src.append((str(36),'Expanded listing 132513280; 4+4; optional garage excluded; later re-fetch failed; factual snapshot, 13 October 2026.','https://www.immobiliare.it/annunci/132513280/'))
from docx.opc.constants import RELATIONSHIP_TYPE as RT
src.append((str(37),'Assaeroporti. 2019 December traffic tables, January–December cumulative total for Bergamo. Summary and passenger-detail totals reconcile. Retrieved 13 October 2026.','https://assaeroporti.com/statistiche/?month=12&selectedYear=2019'))
src.append((str(38),'Assaeroporti. 2020 December traffic tables, January–December cumulative total for Bergamo. Summary and passenger-detail totals reconcile. Retrieved 13 October 2026.','https://assaeroporti.com/statistiche/?month=12&selectedYear=2020'))
src.append((str(39),'Assaeroporti. 2021 December traffic tables, January–December cumulative total for Bergamo. Summary and passenger-detail totals reconcile. Retrieved 13 October 2026.','https://assaeroporti.com/statistiche/?month=12&selectedYear=2021'))
src.append((str(40),'Assaeroporti. 2022 December traffic tables, January–December cumulative total for Bergamo. Summary and passenger-detail totals reconcile. Retrieved 13 October 2026.','https://assaeroporti.com/statistiche/?month=12&selectedYear=2022'))
src.append((str(41),'Assaeroporti. 2023 December traffic tables, January–December cumulative total for Bergamo. Summary and passenger-detail totals reconcile. Retrieved 13 October 2026.','https://assaeroporti.com/statistiche/?month=12&selectedYear=2023'))
src.append((str(42),'Assaeroporti. 2024 December traffic tables, January–December cumulative total for Bergamo. Summary and passenger-detail totals reconcile. Retrieved 13 October 2026.','https://assaeroporti.com/statistiche/?month=12&selectedYear=2024'))
src.append((str(43),'Assaeroporti. 2025 December traffic tables, January–December cumulative total for Bergamo. Summary and passenger-detail totals reconcile. Retrieved 13 October 2026.','https://assaeroporti.com/statistiche/?month=12&selectedYear=2025'))
for num,label,url in src:
    refp=p('['+num+'] '+label);refp.paragraph_format.keep_with_next=True;x=p('');link=OxmlElement('w:hyperlink');link.set(qn('r:id'),x.part.relate_to(url,RT.HYPERLINK,is_external=True));rr=OxmlElement('w:r');rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');rp.append(sz);rr.append(rp);tt=OxmlElement('w:t');tt.text=url;rr.append(tt);link.append(rr);x._p.append(link)
p('Reproducible analysis: https://github.com/pedramkj/bergamo-rental-strategy-analysis. The repository provides Python calculations, validation checks, data snapshots, the listing registry and a README. No personal documents or full copyrighted listing text are included.')
p('Project calculations are retained in validated-monthly-analysis.json, validation.json, seasonal-booking-estimate.json, monthly-financial-scenario.json and cash-breakeven.json. The present model uses exactly 180 nights, replacing the earlier fractional 181.5-night illustration. No owner-time deduction is included. The eight charts are generated from those observations and explicit assumptions.')
for b in list(d.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
d.save(out/'Bergamo_Rental_Strategies_Complete_EN_v03.docx')
print((out/'Bergamo_Rental_Strategies_Complete_EN_v03.docx').resolve())

